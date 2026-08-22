#!/usr/bin/env python3
"""Undo a reprice book: put back exactly the bids and states its executable rows changed.

WHY THIS EXISTS
    The reprice book (build_reprice_bulksheet.py) moves bids and pauses losers on the strength
    of settled evidence. "Reversible" is only a property if the reversing row exists — this is
    that row, learned the hard way (see build_restore_paused_bulksheet.py).

    It reads the AUDIT CSV the book run wrote, so the restore is scoped to exactly what that
    book changed: bid rows get their OLD bid back, pause rows get an enabled row. It never
    re-derives a population — re-deriving would find whatever the states say today, which is
    not the same set and not the question.

    Portfolio is re-read LIVE (the audit stores what was echoed at book time; a campaign whose
    portfolio moved since must not be detached by the restore).

USAGE
    /usr/bin/python3 tools/build_restore_reprice_bulksheet.py --audit .tmp/reprice_book_..._audit.csv
"""
import argparse
import csv
import json
import os
import subprocess
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_stop_nonconverting_bulksheet import (  # noqa: E402
    SP_HEADERS, SB_HEADERS, SP_SHEET, SB_SHEET)

import openpyxl  # noqa: E402

PROJECT = 'onyga-482313'


def bq(sql):
    out = subprocess.run(
        ['bq', 'query', f'--project_id={PROJECT}', '--use_legacy_sql=false',
         '--nouse_cache', '--format=json', '--max_rows=100000', sql],
        capture_output=True, text=True, check=True).stdout
    return json.loads(out or '[]')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--audit', required=True, help='the _audit.csv the book run wrote')
    ap.add_argument('-o', '--out', default=f".tmp/restore_reprice_{date.today():%Y%m%d}.xlsx")
    args = ap.parse_args()

    with open(args.audit) as f:
        rows = [r for r in csv.DictReader(f)
                if r['disposition'] in ('BID_DOWN', 'BID_UP', 'PAUSE')]
    if not rows:
        print('Audit holds no executable rows — nothing to restore.')
        return

    ids = sorted({r['campaign_id'] for r in rows})
    id_list = ','.join(f"'{i}'" for i in ids)
    last_pf = {r['campaign_id']: r['portfolio_id'] for r in bq(f"""
        SELECT CAST(campaign_id AS STRING) AS campaign_id, portfolio_id FROM (
          SELECT campaign_id, portfolio_id,
                 ROW_NUMBER() OVER (PARTITION BY campaign_id ORDER BY date DESC) rn
          FROM `{PROJECT}.OI.V_SRC_AmazonAds_campaign_history`
          WHERE CAST(campaign_id AS STRING) IN ({id_list}) AND portfolio_id IS NOT NULL)
        WHERE rn = 1""")}

    def sp_row(r):
        row = {h: '' for h in SP_HEADERS}
        is_target = r['is_auto'] == 'true' or r['is_pt'] == 'true'
        row.update({
            'Product': 'Sponsored Products',
            'Entity': 'Product Targeting' if is_target else 'Keyword',
            'Operation': 'Update',
            'Campaign ID': r['campaign_id'],
            'Ad Group ID': r['ad_group_id'],
            'Portfolio ID': last_pf.get(r['campaign_id']) or r['portfolio_id'] or '',
            'Campaign Name (Informational only)': r['campaign'],
        })
        key_col = 'Product Targeting ID' if is_target else 'Keyword ID'
        row[key_col] = r['keyword_id']
        if r['disposition'] == 'PAUSE':
            row['State'] = 'ENABLED'
        else:
            row['Bid'] = f"{float(r['old_bid']):.2f}"
        return row

    def sb_row(r):
        row = {h: '' for h in SB_HEADERS}
        row.update({
            'Product': 'Sponsored Brands',
            'Entity': 'Keyword',
            'Operation': 'Update',
            'Campaign Id': r['campaign_id'],
            'Ad Group Id': r['ad_group_id'],
            'Portfolio Id': last_pf.get(r['campaign_id']) or r['portfolio_id'] or '',
            'Keyword Id': r['keyword_id'],
        })
        if r['disposition'] == 'PAUSE':
            row['State'] = 'enabled'
        else:
            row['Bid'] = f"{float(r['old_bid']):.2f}"
        return row

    sp = [r for r in rows if r['channel'].upper() != 'SB']
    sb = [r for r in rows if r['channel'].upper() == 'SB']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if sp:
        ws = wb.create_sheet(SP_SHEET)
        ws.append(SP_HEADERS)
        for r in sp:
            ws.append([sp_row(r)[h] for h in SP_HEADERS])
    if sb:
        ws = wb.create_sheet(SB_SHEET)
        ws.append(SB_HEADERS)
        for r in sb:
            ws.append([sb_row(r)[h] for h in SB_HEADERS])
    os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
    wb.save(args.out)
    print(f"wrote {len(rows)} restore rows ({len(sp)} SP · {len(sb)} SB) -> {args.out}")
    print("Every bid row carries the OLD bid from the audit; every pause row re-enables. "
          "Portfolio read live so the restore cannot detach a campaign.")


if __name__ == '__main__':
    main()
