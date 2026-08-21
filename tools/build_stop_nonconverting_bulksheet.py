#!/usr/bin/env python3
"""Pause the ENABLED campaigns that spent money and produced no ads-attributed sales.

WHY THIS EXISTS
    Periodically a handful of campaigns keep serving while returning nothing. This builds the
    bulksheet that pauses them, with the evidence attached so the decision is reviewable BEFORE
    the upload rather than argued about after it.

WHAT IT EMITS
    One Campaign / Update row per campaign, State = PAUSED (SP) / paused (SB), plus the campaign's
    Portfolio ID echoed back. Nothing else. See sp_row() for why every other mutable column is
    deliberately blank.

    PAUSE, NEVER ARCHIVE. Archiving is irreversible on Amazon: the id is retired and its history,
    portfolio membership and negative keywords go with it. The difference in this format is one
    cell value, which is exactly what makes it dangerous. Most of this population is under-tested
    by the account's own conversion rate (see FLAGS below) and several ids have converted before,
    so the evidence does not support a permanent decision. Pause is reversible with one ENABLED row.

FLAGS (derived every run, never hardcoded)
    UNDER_TESTED      A dry run this short happens by luck at the account's own base rate. At
                      P clicks-per-order, P(zero orders | n clicks) = (1 - 1/P)^n. Flagged when
                      that probability clears LUCK_THRESHOLD, i.e. we cannot rule out chance.
    PRIOR_CONVERTER   The campaign_id has orders in its lifetime history. The 2026-07-25/26
                      intent-grid rename wave re-pointed several campaigns and FACT_AMAZON_ADS
                      keeps the OLD name on the old rows, so grouping by campaign_name shows a
                      virgin campaign while the ID has real history. Always judge at ID grain.
    BRAND_DEFENSE     The campaign buys the brand's own name. House doctrine: brand defense is
                      NEVER judged on ROAS — its payoff is the sale a competitor does NOT take,
                      which can never appear in Ads_sales. Shipped so it is visible, recommended
                      for deletion.
    PORTFOLIO_DETACHED  The live campaign_history row carries portfolio_id NULL (collateral from
                      the earlier blank-Portfolio-ID incident). The row echoes the last non-null
                      portfolio back, so it both pauses AND re-attaches. Deliberate, but flagged
                      so nobody is surprised by a second change riding on a pause row.

USAGE
    /usr/bin/python3 tools/build_stop_nonconverting_bulksheet.py [-o PATH] [--days 28]

    Prints a review table, writes the XLSX, a sibling _audit.csv and a plain-English README.
    PREPARE-ONLY. Nothing here touches Amazon; the upload is manual.

    Safe to re-run: the population is re-derived from the warehouse watermark on every run. There
    is no embedded campaign list, and a campaign already paused drops out on its own.
"""

import argparse
import csv
import json
import re
import subprocess
import sys
from datetime import date, datetime, timezone

import openpyxl

PROJECT = "onyga-482313"

# Window length in complete days. The END of the window is derived from the warehouse watermark at
# run time, never from today's date -- a re-run next week must measure a fresh window, not a stale one.
DEFAULT_DAYS = 28

# A dry run is "not a verdict" when it is at least this likely to have happened by pure chance at
# the account's own base conversion rate. 0.20 => we keep the flag unless we are ~80% confident.
LUCK_THRESHOLD = 0.20

# The brand's own name, for detecting genuine brand-defense traffic. Declared constant.
BRAND_TOKENS = ('happy lolli', 'happylolli', 'happy+lolli')
BRAND_DEFENSE_NAME_RE = re.compile(r'brand\s*defense', re.I)
# Share of clicked traffic on brand terms above which a campaign is treated as brand defense.
BRAND_DEFENSE_CLICK_SHARE = 0.50

# Mirrors SP_HEADERS in dashboard-react/src/pages/DoPage.tsx, and a byte-exact prefix of the header
# row Amazon writes in its own bulk downloads.
SP_HEADERS = [
    'Product', 'Entity', 'Operation',
    'Campaign ID', 'Ad Group ID', 'Portfolio ID', 'Ad ID', 'Keyword ID', 'Product Targeting ID',
    'Campaign Name', 'Ad Group Name',
    'Campaign Name (Informational only)', 'Ad Group Name (Informational only)',
    'Portfolio Name (Informational only)',
    'Start Date', 'End Date', 'Targeting Type', 'State',
    'Campaign State (Informational only)', 'Ad Group State (Informational only)',
    'Daily Budget', 'SKU', 'ASIN (Informational only)',
    'Eligibility Status (Informational only)', 'Reason for Ineligibility (Informational only)',
    'Ad Group Default Bid', 'Ad Group Default Bid (Informational only)',
    'Bid', 'Keyword Text', 'Native Language Keyword', 'Native Language Locale',
    'Match Type', 'Bidding Strategy', 'Placement', 'Percentage',
    'Product Targeting Expression', 'Resolved Product Targeting Expression (Informational only)',
]

# Mirrors SB_HEADERS in DoPage.tsx. NOTE the casing differs from SP ('Campaign Id' not 'Campaign ID')
# and from Amazon's own download; this is the casing the account has actually shipped successfully.
SB_HEADERS = [
    'Product', 'Entity', 'Operation',
    'Campaign Id', 'Ad Group Id', 'Ad Id', 'Keyword Id', 'Draft Campaign Id', 'Portfolio Id',
    'Campaign Name', 'Ad Group Name', 'Ad Name', 'Start Date', 'End Date', 'State',
    'Budget Type', 'Budget',
    'Bid Optimization', 'Bid Multiplier',
    'Bid', 'Keyword Text', 'Match Type',
    'Product Targeting Expression',
    'Ad Format', 'Landing Page URL', 'Landing page ASINs',
    'Brand Entity Id', 'Brand Name',
    'Creative Headline', 'Creative ASINs', 'Video asset IDs', 'Creative Type',
]

SP_SHEET = 'Sponsored Products Campaigns'
SB_SHEET = 'SB Multi Ad Group Campaigns'

SQL = """
WITH wm AS (
  SELECT DATE_SUB(MAX(date), INTERVAL 1 DAY)      AS end_d,
         DATE_SUB(MAX(date), INTERVAL {days} DAY) AS start_d,
         MAX(date)                                AS watermark
  FROM `{p}.OI.FACT_AMAZON_ADS`
),
perf AS (
  SELECT f.campaign_id,
         SUM(f.Ads_cost)   AS spend,
         SUM(f.Ads_clicks) AS clicks,
         SUM(f.Ads_orders) AS orders,
         SUM(f.Ads_sales)  AS sales
  FROM `{p}.OI.FACT_AMAZON_ADS` f, wm
  WHERE f.date BETWEEN wm.start_d AND wm.end_d
  GROUP BY 1
),
-- Account base rate over the same window: how many clicks the account needs to buy one order.
base AS (
  SELECT SAFE_DIVIDE(SUM(f.Ads_clicks), SUM(f.Ads_orders)) AS clicks_per_order
  FROM `{p}.OI.FACT_AMAZON_ADS` f, wm
  WHERE f.date BETWEEN wm.start_d AND wm.end_d
),
-- Live state of each campaign.
latest AS (
  SELECT campaign_id, campaign_name, campaign_type, portfolio_id, state, daily_budget
  FROM `{p}.OI.V_SRC_AmazonAds_campaign_history`
  WHERE campaign_name IS NOT NULL
  QUALIFY ROW_NUMBER() OVER (PARTITION BY campaign_id ORDER BY date DESC) = 1
),
-- Last NON-NULL portfolio the campaign ever had. A blank Portfolio ID on a Campaign Update row
-- means "no portfolio" and DETACHES the campaign, so every row must carry one.
restore AS (
  SELECT campaign_id,
         ARRAY_AGG(portfolio_id IGNORE NULLS ORDER BY date DESC LIMIT 1)[SAFE_OFFSET(0)] AS portfolio_id
  FROM `{p}.OI.V_SRC_AmazonAds_campaign_history`
  GROUP BY campaign_id
),
pf AS (
  SELECT portfolio_id, portfolio_name
  FROM `{p}.OI.V_SRC_AmazonAds_portfolio`
  QUALIFY ROW_NUMBER() OVER (PARTITION BY portfolio_id ORDER BY last_updated_date DESC) = 1
),
-- Lifetime history at CAMPAIGN_ID grain, which survives the rename waves that campaign_name does not.
life AS (
  SELECT campaign_id,
         SUM(Ads_clicks) AS life_clicks,
         SUM(Ads_orders) AS life_orders,
         MAX(IF(Ads_orders > 0, date, NULL)) AS last_order_date,
         STRING_AGG(DISTINCT campaign_name, ' | ') AS known_names
  FROM `{p}.OI.FACT_AMAZON_ADS`
  GROUP BY 1
),
-- What the campaign actually bought in the window, and how much of it was the brand's own name.
terms AS (
  SELECT f.campaign_id,
         STRING_AGG(DISTINCT f.search_term, ' ~ ' ORDER BY f.search_term LIMIT 5) AS sample_terms,
         SAFE_DIVIDE(
           SUM(IF({brand_pred}, f.Ads_clicks, 0)),
           NULLIF(SUM(f.Ads_clicks), 0)
         ) AS brand_click_share
  FROM `{p}.OI.FACT_AMAZON_ADS` f, wm
  WHERE f.date BETWEEN wm.start_d AND wm.end_d AND f.Ads_clicks > 0
  GROUP BY 1
),
-- Recent burn rate: is this campaign still actually spending, or has the bid engine already
-- throttled it to nothing? Changes what "monthly spend stopped" honestly means.
recent AS (
  SELECT f.campaign_id, SUM(f.Ads_cost) AS spend_7d
  FROM `{p}.OI.FACT_AMAZON_ADS` f, wm
  WHERE f.date BETWEEN DATE_SUB(wm.end_d, INTERVAL 6 DAY) AND wm.end_d
  GROUP BY 1
)
SELECT
  CAST(wm.start_d AS STRING)   AS window_start,
  CAST(wm.end_d AS STRING)     AS window_end,
  CAST(wm.watermark AS STRING) AS watermark,
  base.clicks_per_order,
  l.campaign_id, l.campaign_name, l.campaign_type, l.state,
  l.daily_budget,
  l.portfolio_id                AS live_portfolio_id,
  r.portfolio_id                AS echo_portfolio_id,
  pf.portfolio_name,
  p.spend, p.clicks, p.orders, p.sales,
  COALESCE(rc.spend_7d, 0)      AS spend_7d,
  li.life_clicks, li.life_orders,
  CAST(li.last_order_date AS STRING) AS last_order_date,
  li.known_names,
  t.sample_terms, t.brand_click_share
FROM perf p
JOIN latest l USING (campaign_id)
LEFT JOIN restore r USING (campaign_id)
LEFT JOIN pf ON pf.portfolio_id = r.portfolio_id
LEFT JOIN life li ON li.campaign_id = l.campaign_id
LEFT JOIN terms t ON t.campaign_id = l.campaign_id
LEFT JOIN recent rc ON rc.campaign_id = l.campaign_id
CROSS JOIN wm
CROSS JOIN base
WHERE UPPER(COALESCE(l.state, '')) = 'ENABLED'
  AND p.spend > 0
  AND COALESCE(p.sales, 0) = 0     -- zero ads-attributed sales, not merely zero same-ASIN sales
ORDER BY p.spend DESC
"""


def bq(sql):
    out = subprocess.run(
        ['bq', 'query', '--use_legacy_sql=false', '--format=json', '--max_rows=100000',
         '--nouse_cache', f'--project_id={PROJECT}', sql],
        capture_output=True, text=True,
    )
    if out.returncode != 0:
        sys.exit(f"BigQuery failed:\n{out.stderr}")
    return json.loads(out.stdout or '[]')


def fetch(days):
    brand_pred = ' OR '.join(
        f"STRPOS(LOWER(COALESCE(f.search_term,'')), '{t}') > 0" for t in BRAND_TOKENS
    )
    return bq(SQL.format(p=PROJECT, days=days, brand_pred=brand_pred))


def num(v, default=0.0):
    return default if v in (None, '') else float(v)


def luck(clicks, clicks_per_order):
    """P(zero orders in `clicks` clicks | account base rate). High = the dry run proves nothing."""
    if not clicks_per_order or clicks_per_order <= 1:
        return 1.0
    return (1.0 - 1.0 / clicks_per_order) ** clicks


def classify(r):
    """Derive the flags for one campaign. Returns (flags, recommendation)."""
    flags = []
    cpo = num(r.get('clicks_per_order'))
    clicks = int(num(r.get('clicks')))
    p_luck = luck(clicks, cpo)

    if p_luck >= LUCK_THRESHOLD:
        flags.append(('UNDER_TESTED',
                      f"{clicks} clicks with no order happens by pure luck {p_luck:.0%} of the time "
                      f"at the account's own rate of {cpo:.1f} clicks per order — not a verdict"))

    life_orders = int(num(r.get('life_orders')))
    if life_orders > 0:
        flags.append(('PRIOR_CONVERTER',
                      f"this campaign ID has {life_orders} lifetime order(s), last on "
                      f"{r.get('last_order_date')}, across {len(str(r.get('known_names') or '').split(' | '))} "
                      f"name(s) — it has worked before"))

    share = num(r.get('brand_click_share'))
    if BRAND_DEFENSE_NAME_RE.search(r.get('campaign_name') or '') or share >= BRAND_DEFENSE_CLICK_SHARE:
        flags.append(('BRAND_DEFENSE',
                      f"{share:.0%} of its clicks are on the brand's own name — brand defense is "
                      f"never judged on ROAS; its payoff is the sale a competitor does not take"))

    if not r.get('live_portfolio_id') and r.get('echo_portfolio_id'):
        flags.append(('PORTFOLIO_DETACHED',
                      f"currently detached from its portfolio; this row re-attaches it to "
                      f"{r.get('portfolio_name') or r.get('echo_portfolio_id')} as well as pausing it"))

    codes = {c for c, _ in flags}
    if 'BRAND_DEFENSE' in codes:
        rec = 'DELETE THIS LINE unless you have decided to stop defending the brand term'
    elif codes:
        rec = 'INCLUDE — flagged, read the reason before uploading'
    else:
        rec = 'INCLUDE — clean stop, no flag'
    return flags, rec


# Rows carry State and the Portfolio, and NOTHING else mutable -- deliberately.
#
# Portfolio ID must ALWAYS be echoed: blank means "no portfolio" on a Campaign row and detaches the
# campaign (build_portfolio_repair_bulksheet.py exists solely to repair 51 campaigns detached that
# way). Daily Budget, Start Date, Targeting Type and Bidding Strategy are left BLANK, because for
# those columns blank genuinely does mean "unchanged", and restating them creates a REVERT risk:
# this sheet reads the warehouse now and may be uploaded hours later, silently undoing anything
# changed in between. Portfolio is the exception, not the rule.
def sp_row(r):
    row = {h: '' for h in SP_HEADERS}
    row.update({
        'Product': 'Sponsored Products',
        'Entity': 'Campaign',
        'Operation': 'Update',
        'Campaign ID': str(r['campaign_id']),
        'Portfolio ID': str(r['echo_portfolio_id'] or ''),
        # Amazon never writes this column back; it is the human-readable label for review.
        'Campaign Name (Informational only)': r['campaign_name'],
        'State': 'PAUSED',
    })
    return row


def sb_row(r):
    row = {h: '' for h in SB_HEADERS}
    row.update({
        'Product': 'Sponsored Brands',
        'Entity': 'Campaign',
        'Operation': 'Update',
        'Campaign Id': str(r['campaign_id']),
        'Portfolio Id': str(r['echo_portfolio_id'] or ''),
        # 'Campaign Name' on the SB sheet is a REAL column, not informational -- filling it would
        # rename the campaign to whatever the warehouse held at build time. Left blank on purpose;
        # the README maps every row number to its campaign for review.
        'State': 'paused',
    })
    return row


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('-o', '--out', default=f".tmp/stop_nonconverting_{date.today():%Y%m%d}.xlsx")
    ap.add_argument('--days', type=int, default=DEFAULT_DAYS,
                    help=f'complete days in the measurement window (default {DEFAULT_DAYS})')
    args = ap.parse_args()

    rows = fetch(args.days)
    if not rows:
        print("No enabled campaigns spent money without converting — nothing to stop.")
        return

    w0, w1 = rows[0]['window_start'], rows[0]['window_end']
    cpo = num(rows[0].get('clicks_per_order'))
    print(f"Window {w0}..{w1} ({args.days} complete days, watermark {rows[0]['watermark']})")
    print(f"Account base rate: {cpo:.2f} clicks per order\n")

    for r in rows:
        r['_flags'], r['_rec'] = classify(r)
        if not r.get('echo_portfolio_id'):
            print(f"  !! {r['campaign_name']} has NO portfolio anywhere in its history — "
                  f"a blank Portfolio ID would be correct for it, but verify before uploading.")

    sp = [r for r in rows if (r.get('campaign_type') or '').upper() == 'SP']
    sb = [r for r in rows if (r.get('campaign_type') or '').upper() != 'SP']

    total_spend = sum(num(r['spend']) for r in rows)
    total_7d = sum(num(r['spend_7d']) for r in rows)

    print(f"{len(rows)} campaigns to pause  ({len(sp)} SP · {len(sb)} SB)\n")
    hdr = f"{'campaign':<48} {'type':<4} {'spend':>8} {'clk':>4} {'ord':>4} {'7d':>6}  flags"
    print(hdr)
    print('-' * len(hdr))
    for r in rows:
        codes = ','.join(c for c, _ in r['_flags']) or '-'
        print(f"{(r['campaign_name'] or '')[:47]:<48} {r['campaign_type']:<4} "
              f"{num(r['spend']):>8.2f} {int(num(r['clicks'])):>4} {int(num(r['orders'])):>4} "
              f"{num(r['spend_7d']):>6.2f}  {codes}")
    print('-' * len(hdr))
    print(f"{'TOTAL':<48} {'':<4} {total_spend:>8.2f} "
          f"{int(sum(num(r['clicks']) for r in rows)):>4} {0:>4} {total_7d:>6.2f}")

    # ---- workbook -------------------------------------------------------------------
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    line_of = {}
    if sp:
        ws = wb.create_sheet(SP_SHEET)
        ws.append(SP_HEADERS)
        for i, r in enumerate(sp, start=2):
            ws.append([sp_row(r)[h] for h in SP_HEADERS])
            line_of[r['campaign_id']] = (SP_SHEET, i)
    if sb:
        ws = wb.create_sheet(SB_SHEET)
        ws.append(SB_HEADERS)
        for i, r in enumerate(sb, start=2):
            ws.append([sb_row(r)[h] for h in SB_HEADERS])
            line_of[r['campaign_id']] = (SB_SHEET, i)
    wb.save(args.out)

    # ---- audit csv ------------------------------------------------------------------
    audit_path = args.out.rsplit('.', 1)[0] + '_audit.csv'
    with open(audit_path, 'w', newline='') as f:
        wr = csv.writer(f)
        wr.writerow(['sheet', 'excel_row', 'campaign_id', 'campaign', 'type', 'portfolio',
                     'daily_budget', 'spend_window', 'clicks', 'orders', 'spend_last_7d',
                     'life_clicks', 'life_orders', 'last_order', 'p_zero_by_luck',
                     'flags', 'recommendation', 'sample_search_terms'])
        for r in rows:
            sheet, ln = line_of[r['campaign_id']]
            wr.writerow([sheet, ln, r['campaign_id'], r['campaign_name'], r['campaign_type'],
                         r.get('portfolio_name') or r.get('echo_portfolio_id'),
                         r.get('daily_budget'), f"{num(r['spend']):.2f}",
                         int(num(r['clicks'])), int(num(r['orders'])), f"{num(r['spend_7d']):.2f}",
                         int(num(r.get('life_clicks'))), int(num(r.get('life_orders'))),
                         r.get('last_order_date') or 'never',
                         f"{luck(int(num(r['clicks'])), cpo):.2f}",
                         ';'.join(c for c, _ in r['_flags']), r['_rec'],
                         r.get('sample_terms') or ''])

    # ---- plain-English readme -------------------------------------------------------
    readme_path = args.out.rsplit('.', 1)[0] + '_README.md'
    with open(readme_path, 'w') as f:
        f.write(f"# Stop the campaigns that did not convert\n\n")
        f.write(f"Built {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} from "
                f"`{args.out}`.\n\n")
        f.write(f"**Window:** {w0} to {w1} ({args.days} complete days). "
                f"**Account base rate:** {cpo:.1f} clicks per order.\n\n")
        f.write(f"Every row sets the campaign to **paused**. Nothing is archived — pausing is "
                f"reversible with a single ENABLED row, archiving never is.\n\n")
        f.write(f"Each row also re-states the campaign's portfolio. That is mandatory: a blank "
                f"portfolio cell would detach the campaign.\n\n")
        f.write(f"No budget, bid, start date or targeting setting is touched.\n\n")
        f.write(f"**Total spent by these campaigns in the window: ${total_spend:,.2f}.** "
                f"In the most recent 7 days they spent ${total_7d:,.2f} between them, so that "
                f"lower figure is closer to what pausing actually saves going forward.\n\n")
        f.write("---\n\n## Row by row\n\n")
        for r in rows:
            sheet, ln = line_of[r['campaign_id']]
            f.write(f"### {sheet} — row {ln}: {r['campaign_name']}\n\n")
            f.write(f"- **What this row does:** pauses campaign `{r['campaign_id']}` "
                    f"({r['campaign_type']}) and keeps it in portfolio "
                    f"{r.get('portfolio_name') or r.get('echo_portfolio_id')}.\n")
            f.write(f"- **Spent:** ${num(r['spend']):,.2f} · **Clicks:** {int(num(r['clicks']))} "
                    f"· **Orders:** {int(num(r['orders']))} · **Daily budget:** "
                    f"${num(r.get('daily_budget')):,.2f}\n")
            f.write(f"- **Why it is on the list:** it was enabled, it spent money, and it "
                    f"produced no ads-attributed sales at all over the {args.days} days.\n")
            if r['_flags']:
                f.write(f"- **Flags:**\n")
                for code, why in r['_flags']:
                    f.write(f"  - `{code}` — {why}\n")
            f.write(f"- **Recommendation:** {r['_rec']}\n")
            f.write(f"- **Delete this line and** campaign `{r['campaign_id']}` keeps running "
                    f"exactly as it is today; nothing else in the sheet changes.\n\n")

    stamp = datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
    print(f"\n[{stamp}] wrote {len(rows)} pause rows -> {args.out}")
    print(f"  audit  -> {audit_path}")
    print(f"  readme -> {readme_path}")
    print(f"  window spend ${total_spend:,.2f} · last-7d spend ${total_7d:,.2f}")
    print("  PREPARE-ONLY: review the README, delete any line you disagree with, "
          "then upload to Amazon Ads > Bulk operations manually.")


if __name__ == '__main__':
    main()
