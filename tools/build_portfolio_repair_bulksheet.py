#!/usr/bin/env python3
"""Re-attach campaigns that lost their portfolio to a blank bulksheet column.

WHY THIS EXISTS
    A Campaign Update row with an empty Portfolio ID does NOT mean "leave the portfolio alone".
    Amazon reads it as "no portfolio" and detaches the campaign. The DO-page exporter shipped every
    budget row that way (spBase never set the column; the writer filled missing headers with ''),
    so each budget upload quietly stripped the portfolio off every campaign it touched — 16 of them
    at the 2026-08-09 12:56 upload alone, 51 detached in total.

    The exporter is fixed (DoPage.tsx echoes the portfolio back on Campaign rows, sourced from
    /api/live-campaigns). This script repairs the campaigns already detached.

WHAT IT EMITS
    One Campaign Update row per detached campaign, carrying the LAST NON-NULL portfolio seen in
    campaign_history — and no other mutable field, so the sheet cannot revert a budget or a pause
    applied between building it and uploading it (see the note above sp_row).

    SP campaigns -> "Sponsored Products Campaigns" sheet; SB -> "SB Multi Ad Group Campaigns".

USAGE
    python3 tools/build_portfolio_repair_bulksheet.py [-o .tmp/portfolio_repair_<date>.xlsx]

    Prints the repair table, writes the XLSX. Upload is manual — nothing here touches Amazon.
    Safe to re-run: a campaign whose portfolio is already intact is not emitted.
"""

import argparse
import json
import subprocess
import sys
from datetime import date

import openpyxl

PROJECT = "onyga-482313"

# Mirrors SP_HEADERS / SB_HEADERS in dashboard-react/src/pages/DoPage.tsx
SP_HEADERS = [
    'Product', 'Entity', 'Operation',
    'Campaign ID', 'Ad Group ID', 'Portfolio ID', 'Ad ID', 'Keyword ID', 'Product Targeting ID',
    'Campaign Name', 'Ad Group Name',
    'Campaign Name (Informational only)', 'Ad Group Name (Informational only)',
    'Portfolio Name (Informational only)',
    'Start Date', 'End Date', 'Targeting Type', 'State',
    'Campaign State (Informational only)', 'Ad Group State (Informational only)',
    'Daily Budget', 'SKU', 'ASIN (Informational only)',
    'Eligibility Status (Informational only)', 'Reason for Ineligibility (Informational only)',
    'Ad Group Default Bid', 'Ad Group Default Bid (Informational only)',
    'Bid', 'Keyword Text', 'Native Language Keyword', 'Native Language Locale',
    'Match Type', 'Bidding Strategy', 'Placement', 'Percentage',
    'Product Targeting Expression', 'Resolved Product Targeting Expression (Informational only)',
]

SB_HEADERS = [
    'Product', 'Entity', 'Operation',
    'Campaign Id', 'Ad Group Id', 'Ad Id', 'Keyword Id', 'Draft Campaign Id', 'Portfolio Id',
    'Campaign Name', 'Ad Group Name', 'Ad Name', 'Start Date', 'End Date', 'State',
    'Budget Type', 'Budget',
    'Bid Optimization', 'Bid Multiplier',
    'Bid', 'Keyword Text', 'Match Type',
    'Product Targeting Expression',
    'Ad Format', 'Landing Page URL', 'Landing page ASINs',
    'Brand Entity Id', 'Brand Name',
    'Creative Headline', 'Creative ASINs', 'Video asset IDs', 'Creative Type',
]

# Campaigns detached from their portfolio, with the portfolio to put them back into.
#   latest  = the campaign's most recent history row (its live state)
#   restore = the most recent NON-NULL portfolio it ever had
# A campaign qualifies only when latest.portfolio_id IS NULL and a restore value exists, so the
# script is idempotent — re-running after a successful upload emits nothing.
SQL = """
WITH latest AS (
  SELECT campaign_id, campaign_name, campaign_type, portfolio_id, state, daily_budget
  FROM `onyga-482313.OI.V_SRC_AmazonAds_campaign_history`
  WHERE campaign_name IS NOT NULL
  QUALIFY ROW_NUMBER() OVER (PARTITION BY campaign_id ORDER BY date DESC) = 1
),
restore AS (
  SELECT campaign_id,
         ARRAY_AGG(portfolio_id IGNORE NULLS ORDER BY date DESC LIMIT 1)[SAFE_OFFSET(0)] AS portfolio_id
  FROM `onyga-482313.OI.V_SRC_AmazonAds_campaign_history`
  GROUP BY campaign_id
)
SELECT l.campaign_id, l.campaign_name, l.campaign_type, l.state, l.daily_budget,
       r.portfolio_id AS restore_portfolio_id,
       p.name         AS restore_portfolio_name
FROM latest l
JOIN restore r USING (campaign_id)
LEFT JOIN (
  SELECT portfolio_id, portfolio_name AS name
  FROM `onyga-482313.OI.V_SRC_AmazonAds_portfolio`
  QUALIFY ROW_NUMBER() OVER (PARTITION BY portfolio_id ORDER BY last_updated_date DESC) = 1
) p ON p.portfolio_id = r.portfolio_id
WHERE l.portfolio_id IS NULL
  AND r.portfolio_id IS NOT NULL
  AND UPPER(COALESCE(l.state, '')) != 'ARCHIVED'
ORDER BY l.campaign_type, l.campaign_name
"""


def bq(sql):
    out = subprocess.run(
        ['bq', 'query', '--use_legacy_sql=false', '--format=json', '--max_rows=100000',
         f'--project_id={PROJECT}', sql],
        capture_output=True, text=True,
    )
    if out.returncode != 0:
        sys.exit(f"BigQuery failed:\n{out.stderr}")
    return json.loads(out.stdout or '[]')


# Rows carry the portfolio and NOTHING else mutable — deliberately.
#
# The tempting move is to restate State and Budget so nothing can be read as "clear". But the very
# uploads that caused this damage prove blank means "unchanged" for those columns: every detached
# campaign kept its State through a bulksheet that left State blank. Portfolio is the exception, not
# the rule. Restating budget/state would instead create a REVERT risk — this sheet reads the warehouse
# now and may be uploaded later, so a budget or pause Ori applies in between would be silently undone.
# The campaign name rides along as "(Informational only)", which Amazon never writes back.
def sp_row(r):
    row = {h: '' for h in SP_HEADERS}
    row.update({
        'Product': 'Sponsored Products',
        'Entity': 'Campaign',
        'Operation': 'Update',
        'Campaign ID': r['campaign_id'],
        'Portfolio ID': r['restore_portfolio_id'],
        'Campaign Name (Informational only)': r['campaign_name'],
    })
    return row


def sb_row(r):
    row = {h: '' for h in SB_HEADERS}
    row.update({
        'Product': 'Sponsored Brands',
        'Entity': 'Campaign',
        'Operation': 'Update',
        'Campaign Id': r['campaign_id'],
        'Portfolio Id': r['restore_portfolio_id'],
    })
    return row


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('-o', '--out', default=f".tmp/portfolio_repair_{date.today():%Y%m%d}.xlsx")
    args = ap.parse_args()

    rows = bq(SQL)
    if not rows:
        print("No detached campaigns — nothing to repair.")
        return

    sp = [r for r in rows if r['campaign_type'] == 'SP']
    sb = [r for r in rows if r['campaign_type'] != 'SP']

    print(f"{len(rows)} campaigns to re-attach  ({len(sp)} SP · {len(sb)} SB)\n")
    print(f"{'campaign':<50} {'state':<8} {'budget':>8}  portfolio")
    print("-" * 100)
    for r in rows:
        pf = r['restore_portfolio_name'] or r['restore_portfolio_id']
        print(f"{r['campaign_name'][:49]:<50} {(r['state'] or ''):<8} "
              f"{str(r['daily_budget'] or ''):>8}  {pf}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if sp:
        ws = wb.create_sheet('Sponsored Products Campaigns')
        ws.append(SP_HEADERS)
        for r in sp:
            ws.append([sp_row(r)[h] for h in SP_HEADERS])
    if sb:
        ws = wb.create_sheet('SB Multi Ad Group Campaigns')
        ws.append(SB_HEADERS)
        for r in sb:
            ws.append([sb_row(r)[h] for h in SB_HEADERS])
    wb.save(args.out)
    print(f"\nWrote {args.out} — upload manually to Amazon Ads > Bulk operations.")


if __name__ == '__main__':
    main()
