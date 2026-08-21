#!/usr/bin/env python3
"""Undo a pause bulksheet: re-enable exactly the campaigns a stop sheet paused.

WHY THIS EXISTS
    build_stop_nonconverting_bulksheet.py deliberately PAUSES rather than archives, on the grounds
    that pause "is reversible with one ENABLED row". That is only true if the ENABLED row actually
    exists somewhere — otherwise reversibility is a claim, not a property. This is that row.

    It reads the AUDIT CSV the stop run wrote, so the restore is scoped to exactly what was paused
    on that date. It never re-derives a population: re-deriving would find whatever looks paused
    today, which is not the same set and not the question.

WHAT IT EMITS
    One Campaign / Update row per campaign, State = ENABLED (SP) / enabled (SB), with the campaign's
    portfolio echoed back from the warehouse. Identical column layout and sheet routing to the stop
    sheet, imported from it so the two can never drift apart.

    THE PORTFOLIO ECHO IS NOT COSMETIC. A blank Portfolio ID on a Campaign Update row DETACHES the
    campaign; blank does not mean "unchanged". The last non-null portfolio is read live rather than
    taken from the audit CSV, which stores the portfolio NAME.

USAGE
    /usr/bin/python3 tools/build_restore_paused_bulksheet.py --audit .tmp/stop_..._audit.csv
"""
import argparse, csv, subprocess, json, sys, os
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_stop_nonconverting_bulksheet import (  # noqa: E402
    SP_HEADERS, SB_HEADERS, SP_SHEET, SB_SHEET)

import openpyxl  # noqa: E402

PROJECT = 'onyga-482313'


def bq(sql):
    out = subprocess.run(
        ['bq', 'query', f'--project_id={PROJECT}', '--use_legacy_sql=false',
         '--nouse_cache', '--format=json', sql],
        capture_output=True, text=True, check=True).stdout
    return json.loads(out or '[]')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--audit', required=True, help='the _audit.csv the stop run wrote')
    ap.add_argument('-o', '--out', default=f".tmp/restore_paused_{date.today():%Y%m%d}.xlsx")
    args = ap.parse_args()

    with open(args.audit) as f:
        paused = list(csv.DictReader(f))
    if not paused:
        print('Audit file holds no rows — nothing to restore.')
        return
    ids = [r['campaign_id'] for r in paused]

    # Portfolio and type read LIVE. The audit stores the portfolio NAME; Amazon needs the ID, and a
    # blank one detaches the campaign.
    id_list = ','.join(f"'{i}'" for i in ids)
    live = {r['campaign_id']: r for r in bq(f"""
        SELECT CAST(campaign_id AS STRING) AS campaign_id,
               ANY_VALUE(campaign_name) AS campaign_name,
               ANY_VALUE(campaign_type) AS campaign_type,
               ANY_VALUE(campaign_state) AS campaign_state,
               ANY_VALUE(portfolio_id)  AS portfolio_id
        FROM `{PROJECT}.OI.V_DIM_CAMPAIGN_CURRENT`
        WHERE CAST(campaign_id AS STRING) IN ({id_list}) GROUP BY 1""")}

    # A campaign whose live portfolio has gone NULL keeps its last non-null one, same rule the stop
    # sheet used — otherwise the restore row would detach what the pause row re-attached.
    last_pf = {r['campaign_id']: r['portfolio_id'] for r in bq(f"""
        SELECT CAST(campaign_id AS STRING) AS campaign_id, portfolio_id FROM (
          SELECT campaign_id, portfolio_id,
                 ROW_NUMBER() OVER (PARTITION BY campaign_id ORDER BY date DESC) rn
          FROM `{PROJECT}.OI.V_SRC_AmazonAds_campaign_history`
          WHERE CAST(campaign_id AS STRING) IN ({id_list}) AND portfolio_id IS NOT NULL)
        WHERE rn = 1""")}

    rows = []
    for p in paused:
        cid = p['campaign_id']
        lv = live.get(cid, {})
        rows.append({
            'campaign_id': cid,
            'campaign_name': lv.get('campaign_name') or p['campaign'],
            'campaign_type': (lv.get('campaign_type') or p['type']).upper(),
            'echo_portfolio_id': lv.get('portfolio_id') or last_pf.get(cid) or '',
            'state_now': lv.get('campaign_state') or '?',
            'flags': p.get('flags', ''),
        })

    def sp_row(r):
        row = {h: '' for h in SP_HEADERS}
        row.update({'Product': 'Sponsored Products', 'Entity': 'Campaign', 'Operation': 'Update',
                    'Campaign ID': str(r['campaign_id']),
                    'Portfolio ID': str(r['echo_portfolio_id'] or ''),
                    'Campaign Name (Informational only)': r['campaign_name'],
                    'State': 'ENABLED'})
        return row

    def sb_row(r):
        row = {h: '' for h in SB_HEADERS}
        row.update({'Product': 'Sponsored Brands', 'Entity': 'Campaign', 'Operation': 'Update',
                    'Campaign Id': str(r['campaign_id']),
                    'Portfolio Id': str(r['echo_portfolio_id'] or ''),
                    'Campaign Name': r['campaign_name'], 'State': 'enabled'})
        return row

    sp = [r for r in rows if r['campaign_type'] == 'SP']
    sb = [r for r in rows if r['campaign_type'] != 'SP']

    print(f"{'campaign':<52} {'type':<4} {'portfolio':<18} {'state now':<10} flags")
    print('-' * 110)
    for r in rows:
        print(f"{r['campaign_name'][:50]:<52} {r['campaign_type']:<4} "
              f"{str(r['echo_portfolio_id'])[:16]:<18} {r['state_now']:<10} {r['flags']}")
    missing = [r['campaign_name'] for r in rows if not r['echo_portfolio_id']]
    if missing:
        print(f"\n!! NO PORTFOLIO RESOLVED for {len(missing)}: {missing}. A blank Portfolio ID "
              f"DETACHES the campaign — resolve these by hand before uploading.")

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    if sp:
        ws = wb.create_sheet(SP_SHEET); ws.append(SP_HEADERS)
        for r in sp: ws.append([sp_row(r)[h] for h in SP_HEADERS])
    if sb:
        ws = wb.create_sheet(SB_SHEET); ws.append(SB_HEADERS)
        for r in sb: ws.append([sb_row(r)[h] for h in SB_HEADERS])
    wb.save(args.out)
    print(f"\nWrote {args.out} — {len(sp)} SP row(s), {len(sb)} SB row(s).")
    print("Uploading this re-enables every campaign the stop sheet paused. To restore only some, "
          "delete the other lines first.")


if __name__ == '__main__':
    main()
